"""Monta Supplementary_AAP.docx (Tabelas S1-S14) a partir de resultados/*.csv e dados/."""
import os
import numpy as np, pandas as pd
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

B = os.path.join(os.path.dirname(__file__), '..')
R, D = os.path.join(B, 'resultados'), os.path.join(B, 'dados')
doc = Document()
st = doc.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
sec = doc.sections[0]; sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2)


def para(t, bold=False, size=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph(); p.alignment = align; r = p.add_run(t); r.bold = bold
    if size: r.font.size = Pt(size)
    return p


def table(df, caption, note=None, fmt=None):
    para(caption, bold=True, size=10)
    rows = [list(df.columns)] + df.astype(str).values.tolist()
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for ri, row in enumerate(rows):
        for ci, c in enumerate(row):
            cell = t.cell(ri, ci); cell.paragraphs[0].paragraph_format.line_spacing = 1.0
            run = cell.paragraphs[0].add_run(c); run.font.size = Pt(8.5 if len(rows[0]) > 5 else 9); run.bold = (ri == 0) or None
    tbl = t._tbl; tblPr = tbl.tblPr; b = OxmlElement('w:tblBorders')
    for edge, v in [('top', 'single'), ('bottom', 'single'), ('insideH', 'nil'), ('insideV', 'nil'), ('left', 'nil'), ('right', 'nil')]:
        e = OxmlElement('w:' + edge); e.set(qn('w:val'), v); e.set(qn('w:sz'), '6'); b.append(e)
    tblPr.append(b)
    for c in t.rows[0].cells:
        tcPr = c._tc.get_or_add_tcPr(); bb = OxmlElement('w:tcBorders'); e = OxmlElement('w:bottom'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '6'); bb.append(e); tcPr.append(bb)
    if note:
        para(note, size=9)
    doc.add_paragraph()


def orci(df, dec=2):
    return df.apply(lambda r: f"{r['OR']:.{dec}f} ({r['lo']:.{dec}f}–{r['hi']:.{dec}f})", axis=1)


para("Supplementary material — Hourly rainfall and crash risk on Brazil's federal highways: a time-stratified case-crossover study linking 323,000 crashes to automatic weather stations, 2017–2025", bold=True, size=12)
para('Francisco Lledo dos Santos, Universidade do Estado de Mato Grosso (UNEMAT), ProfÁgua, Cáceres, MT, Brazil', size=10)
doc.add_paragraph()

# ---- S1 harmonização
acc = pd.read_parquet(os.path.join(D, 'prf_acidentes_2017_2025.parquet'), columns=['tipo_acidente'])
TIPO = {'Saída de leito carroçável': 'Run-off-road', 'Capotamento': 'Rollover/overturn', 'Tombamento': 'Rollover/overturn',
        'Colisão com objeto estático': 'Collision with object', 'Colisão com objeto': 'Collision with object', 'Colisão com objeto em movimento': 'Collision with object',
        'Colisão frontal': 'Head-on (incl. opposite-direction sideswipe)', 'Colisão lateral sentido oposto': 'Head-on (incl. opposite-direction sideswipe)',
        'Colisão traseira': 'Rear-end (incl. pile-up)', 'Engavetamento': 'Rear-end (incl. pile-up)',
        'Colisão lateral': 'Sideswipe/angle', 'Colisão lateral mesmo sentido': 'Sideswipe/angle', 'Colisão transversal': 'Sideswipe/angle',
        'Atropelamento de Pedestre': 'Pedestrian struck', 'Atropelamento de Animal': 'Animal struck',
        'Queda de ocupante de veículo': 'Occupant fall (mostly motorcycles)', 'Incêndio': 'Vehicle fire'}
vc = acc['tipo_acidente'].value_counts()
s1 = pd.DataFrame({'PRF crash type (original)': vc.index, 'Harmonized group': [TIPO.get(k, 'Other') for k in vc.index], 'Crashes 2017–2025': vc.values})
table(s1, 'Table S1. Harmonization of PRF crash types (all 632,659 valid records).')

# ---- S2 concordância
c = pd.read_csv(os.path.join(R, 'concordancia_inmet_prf.csv'), index_col=0)
c.columns = ['Officer: no rain', 'Officer: rain or drizzle']
c.index = ['No rain (<0.2 mm)', 'Rain (≥0.2 mm)']
c['% officer rain'] = (100 * c['Officer: rain or drizzle'] / c.sum(1)).round(1)
agree = (c.iloc[0, 0] + c.iloc[1, 1]) / c.iloc[:, :2].values.sum()
s2 = c.reset_index().rename(columns={'index': 'INMET station, crash hour'})
table(s2, f'Table S2. Agreement between station rainfall in the crash hour and the weather condition recorded by the PRF officer (linked cases; overall agreement {100*agree:.1f}%).')

# ---- S3 subgrupos
res = pd.read_csv(os.path.join(R, 'cc_resultados.csv'))
TR = {'classificacao': 'PRF class', 'gravidade': 'Severity', 'regiao': 'Region', 'periodo': 'Light', 'pista': 'Carriageway', 'tipo': 'Crash type', 'area': 'Land use'}
VAL = {'sem vitimas': 'No victims', 'com feridos': 'Injury', 'com mortos': 'Fatal', 'sem vitima grave': 'No serious injury', 'grave nao fatal': 'Serious, non-fatal', 'fatal': 'Fatal', 'noite': 'Night', 'dia': 'Day',
       'transicao': 'Dawn/dusk', 'Simples': 'Single', 'Dupla': 'Dual', 'Múltipla': 'Multilane', 'rural': 'Rural', 'urbana': 'Urban',
       'Colisão traseira': 'Rear-end (incl. pile-up)', 'Saída de pista': 'Run-off-road', 'Capotamento/tombamento': 'Rollover/overturn', 'Colisão com objeto': 'Collision with object', 'Lateral/transversal': 'Sideswipe/angle', 'Atropelamento de pedestre': 'Pedestrian struck', 'Atropelamento de animal': 'Animal struck', 'Queda de ocupante': 'Occupant fall (mostly motorcycles)', 'Incêndio': 'Vehicle fire', 'Outros': 'Other', 'Colisão frontal': 'Head-on (incl. opposite sideswipe)',
       'S': 'South', 'N': 'North', 'SE': 'Southeast', 'NE': 'Northeast', 'CO': 'Centre-West'}
s = res[res['modelo'].str.startswith('S ')].copy()
s['Subgroup'] = s['modelo'].str.extract(r'S (\w+):')[0].map(TR) + ': ' + s['modelo'].str.split(': ').str[1].map(lambda v: VAL.get(v, v))
s3 = pd.DataFrame({'Subgroup': s['Subgroup'], 'OR (95% CI)': orci(s, 3), 'p': s['p'].map(lambda v: f'{v:.1e}' if v < 0.001 else f'{v:.3f}'),
                   'q (BH)': s['q_BH'].map(lambda v: f'{v:.1e}' if v < 0.001 else f'{v:.3f}'), 'Cases': s['n_casos']})
table(s3, 'Table S3. Odds ratios for rain (≥0.2 mm) in the crash hour by subgroup (separate case-crossover models; station-clustered 95% CI; Benjamini–Hochberg q across the subgroup estimates).',
      'The PRF class "Fatal" and the severity class "Fatal" are the same crashes and the same model.')

# ---- S4 anual
ya = res[res['modelo'].str.startswith('Ano ')].copy(); ya['Year'] = ya['modelo'].str[-4:]
comp = pd.read_csv(os.path.join(R, 'composicao_classificacao_ano_pareada.csv'), index_col=0)
s4 = pd.DataFrame({'Year': ya['Year'], 'OR (95% CI)': orci(ya, 3), 'Cases': ya['n_casos'],
                   'Linked crashes without victims (%)': [f"{100*comp.loc[int(y), 'Sem Vítimas']:.1f}" for y in ya['Year']]})
table(s4, 'Table S4. Annual odds ratios for rain (≥0.2 mm) in the crash hour and share of linked crashes without victims.')

# ---- S5 LOYO
lo = pd.read_csv(os.path.join(R, 'preditivo_loyo.csv'))
lo = lo[lo['alvo'] == 'y'].copy()
MN = {'B0': 'GLM', 'B1': 'GLM + rain', 'L0': 'LightGBM', 'L1': 'LightGBM + rain'}
s5 = pd.DataFrame({'Test year': lo['ano_teste'], 'Model': lo['modelo'].map(MN), 'Mean Poisson deviance': lo['desvio'].round(4),
                   'Skill vs GLM (%)': lo['skill_vs_B0_%'].round(3), 'Obs/pred heavy-rain days': lo['obs_prev_chuva_forte'].round(3),
                   'Obs/pred dry days': lo['obs_prev_seco'].round(3), 'Heavy-rain station-days': lo['n_forte']})
table(s5, 'Table S5. Leave-one-year-out results by test year (all crashes). Skill = percentage reduction in mean Poisson deviance relative to the calendar-only GLM in the same test year.')

# ---- S6 defasagens
lag = pd.read_csv(os.path.join(R, 'defasagens.csv'))
cl = pd.read_csv(os.path.join(R, 'rev17_corr_lags.csv'))
s6 = pd.DataFrame({'Lag (h)': lag['lag_h'].astype(int), 'OR (95% CI)': orci(lag, 3)})
table(s6, f"Table S6. Distributed-lag odds ratios for rain (≥0.2 mm) in the crash hour (lag 0) and in each of the six previous hours, estimated jointly ({int(lag['n_casos'][0]):,} crashes with complete lag data).",
      f"Phi correlation between adjacent lag indicators: lag 0–lag 1 = {cl.loc[0, 'phi']:.2f}; lag 1–lag 2 = {cl.loc[1, 'phi']:.2f}.")

# ---- S7 missingness
my = pd.read_csv(os.path.join(R, 'rev17_missing_por_ano.csv'))
s7 = pd.DataFrame({'Year': my['ano'], 'Stations in file': my['estacoes'], 'Missing station-hours (%)': (100 * my['frac_faltante']).round(1)})
table(s7, 'Table S7. Share of missing station-hours in the INMET automatic network, by year (complete hourly grid of all stations in the annual file).')
mc = pd.read_csv(os.path.join(R, 'rev17_missing_condicional.csv'))
mc['Condition'] = mc['condicao'].map({'hora valida seca (<0.2)': 'Valid hour, dry (<0.2 mm)', 'hora valida chuva 0.2-10': 'Valid hour, 0.2–10 mm', 'hora valida chuva >=10': 'Valid hour, ≥10 mm'})
mc = mc.set_index('Condition').loc[['Valid hour, dry (<0.2 mm)', 'Valid hour, 0.2–10 mm', 'Valid hour, ≥10 mm']].reset_index()
s7b = pd.DataFrame({'Rainfall in the valid station-hour': mc['Condition'], 'Station-hours': mc['n_horas'].astype(int),
                    'Next hour missing (%)': (100 * mc['p_falta_seguinte']).round(2), 'Previous hour missing (%)': (100 * mc['p_falta_anterior']).round(2)})
table(s7b, 'Table S8. Probability that the station-hour following or preceding a valid station-hour is missing, by rainfall in the valid hour, 2017–2025 (all stations in the annual files, complete hourly grid).')
rf = pd.read_csv(os.path.join(R, 'rev17_referentes_faltantes_por_exposicao.csv'))
rf['Case-hour rainfall (mm h⁻¹)'] = rf['cat'].map({'0': '<0.2 (dry)', '0.2-2': '0.2–2', '2-5': '2–5', '5-10': '5–10', '>=10': '≥10', 'chuva (>=0.2)': 'Any rain (≥0.2)'})
s7c = pd.DataFrame({'Case-hour rainfall (mm h⁻¹)': rf['Case-hour rainfall (mm h⁻¹)'], 'Cases': rf['casos'].astype(int), 'Candidate referent hours': rf['ref_candidatos'].astype(int),
                    'Valid referent hours': rf['ref_validos'].astype(int), 'Missing referents (%)': (100 * rf['frac_ref_faltante']).round(2)})
table(s7c, 'Table S9. Missing referent hours by rainfall in the case hour (linked strata). Candidate referents are the other same-weekday hours in the calendar month.')

# ---- S10 sensibilidade missingness
sm = pd.read_csv(os.path.join(R, 'rev17_sens_missing.csv'))
m1 = res[res['modelo'] == 'M1 intensidade']
rows = [dict(Analysis='Main analysis (all strata)', Cases=int(m1['n_casos'].iloc[0]), **{t: v for t, v in zip(['0.2–2', '2–5', '5–10', '≥10'], orci(m1))})]
NM = {'Sens sem 2021 e 2022 (anos com >30% de falta)': 'Excluding 2021 and 2022 (years with >30% missing station-hours)',
      'Sens sem estação-anos com >30% de falta': 'Excluding station-years with >30% missing hours',
      'Sens sem estratos com hora vizinha a lacuna': 'Excluding strata with a missing hour adjacent to the case or a referent hour',
      'Sens só estratos com todos os referentes válidos': 'Only strata with all candidate referents valid'}
for m, g in sm.groupby('modelo', sort=False):
    rows.append(dict(Analysis=NM.get(m, m), Cases=int(g['n_casos'].iloc[0]), **{t: v for t, v in zip(['0.2–2', '2–5', '5–10', '≥10'], orci(g))}))
s10 = pd.DataFrame(rows)[['Analysis', 'Cases', '0.2–2', '2–5', '5–10', '≥10']]
table(s10, 'Table S10. Sensitivity of the intensity odds ratios (95% CI) to station data loss. Columns: rainfall in the crash hour (mm h⁻¹), reference <0.2 mm h⁻¹.')

# ---- S11 dose-resposta
dr = pd.read_csv(os.path.join(R, 'rev17_dose_resposta.csv')); dt = pd.read_csv(os.path.join(R, 'rev17_dose_resposta_testes.csv'))
DRN = {('DR log(1+p0) contínuo', 'logp'): 'OR per unit of ln(1 + mm h⁻¹) (continuous)',
       ('DR escore ordinal (por categoria)', 'score'): 'OR per intensity class (ordinal score 0–4, linear trend)',
       ('DR tendência acima de 2 mm (escore 0,1,2 em 2-5,5-10,>=10)', 'p0_0.2-2'): 'Piecewise: 0.2–2 mm h⁻¹ vs dry',
       ('DR tendência acima de 2 mm (escore 0,1,2 em 2-5,5-10,>=10)', 'hi_any'): 'Piecewise: ≥2 mm h⁻¹ vs dry (level)',
       ('DR tendência acima de 2 mm (escore 0,1,2 em 2-5,5-10,>=10)', 's_hi'): 'Piecewise: per class step above 2 mm h⁻¹ (2–5 → 5–10 → ≥10)'}
dr2 = dr[[(m, t) in DRN for m, t in zip(dr['modelo'], dr['termo'])]].copy()
dr2['Term'] = [DRN[(m, t)] for m, t in zip(dr2['modelo'], dr2['termo'])]
s11 = pd.DataFrame({'Term': dr2['Term'], 'OR (95% CI)': orci(dr2, 3), 'p': dr2['p'].map(lambda v: f'{v:.1e}' if v < 0.001 else f'{v:.3f}')})
w = dt.set_index('teste')
table(s11, 'Table S11. Formal tests of the dose–response shape (conditional logistic regression, station-clustered errors, 323,501 crashes).',
      f"Wald tests on the categorical model: equality of the 2–5, 5–10 and ≥10 mm h⁻¹ coefficients (plateau), χ² = {w.loc['plato_2-5_5-10_>=10', 'chi2']:.2f}, 2 df, p = {w.loc['plato_2-5_5-10_>=10', 'p']:.3f}; "
      f"difference between 0.2–2 and 2–5 mm h⁻¹, χ² = {w.loc['degrau_0.2-2_vs_2-5', 'chi2']:.1f}, 1 df, p = {w.loc['degrau_0.2-2_vs_2-5', 'p']:.1e}.")

# ---- S12 distância x uso do solo
dx = pd.read_csv(os.path.join(R, 'rev17_distancia_x_uso_solo.csv'))
dx['Stratum'] = dx['modelo'].map({'rural <=10 km': 'Rural, station ≤10 km', 'rural >10 km': 'Rural, station >10 km', 'urbana <=10 km': 'Urban, station ≤10 km', 'urbana >10 km': 'Urban, station >10 km'})
s12 = pd.DataFrame({'Stratum': dx['Stratum'], 'OR (95% CI)': orci(dx), 'Cases': dx['n_casos']})
table(s12, 'Table S12. Odds ratios for rain (≥0.2 mm) in the crash hour by linkage distance and land use (separate models).')

# ---- S13 OR anual por classe
yc = pd.read_csv(os.path.join(R, 'rev17_OR_anual_por_classe.csv')); tc = pd.read_csv(os.path.join(R, 'rev17_tendencia_por_classe.csv'))
CL = {'com feridos': 'Injury', 'sem vitimas': 'No victims', 'com mortos': 'Fatal'}
piv = yc.pivot(index='ano', columns='classe', values='OR'); lo_ = yc.pivot(index='ano', columns='classe', values='lo'); hi_ = yc.pivot(index='ano', columns='classe', values='hi')
s13 = pd.DataFrame({'Year': piv.index})
for cl_ in ['sem vitimas', 'com feridos', 'com mortos']:
    s13[CL[cl_]] = [f"{piv.loc[y, cl_]:.2f} ({lo_.loc[y, cl_]:.2f}–{hi_.loc[y, cl_]:.2f})" if cl_ in piv.columns and not np.isnan(piv.loc[y, cl_]) else '—' for y in piv.index]
tnote = '; '.join(f"{CL[r['classe']]}, {('all years' if r['anos']=='todos' else '2018–2025')}: {r['tendencia_logOR_ano']:+.4f} log-OR per year (p = {r['p']:.3g})" for _, r in tc.iterrows())
table(s13, 'Table S13. Annual odds ratios (95% CI) for rain (≥0.2 mm) in the crash hour within each PRF class.',
      'Inverse-variance weighted linear trends: ' + tnote + '.')

# ---- S14 atribuíveis
at = pd.read_csv(os.path.join(R, 'rev17_atribuiveis_IC.csv'))
s14 = pd.DataFrame({'Quantity': ['Crashes attributable to rain in the crash hour', 'Attributable fraction among linked crashes (%)'],
                    'Estimate (95% CI)': [f"{at['atribuiveis'][0]:,.0f} ({at['lo'][0]:,.0f}–{at['hi'][0]:,.0f})", f"{at['FA_pct'][0]:.2f} ({at['FA_lo'][0]:.2f}–{at['FA_hi'][0]:.2f})"]})
table(s14, 'Table S14. Attributable crashes with 95% CI from 20,000 draws of the intensity coefficients from their multivariate normal distribution with the station-clustered covariance matrix.')
out = os.path.join(B, 'Supplementary_AAP.docx'); doc.save(out); print('saved', out)
