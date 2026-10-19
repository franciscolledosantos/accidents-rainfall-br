"""Gera Cover_letter_AAP.docx, Highlights_AAP.docx e LEIA-ME_AAP.docx a partir dos .md/.txt."""
import os
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
B = os.path.join(os.path.dirname(__file__), '..')


def mk(src, out, size=12):
    doc = Document(); st = doc.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(size)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    for line in open(os.path.join(B, src), encoding='utf-8').read().split('\n'):
        if not line.strip():
            continue
        p = doc.add_paragraph()
        if line.startswith('#'):
            r = p.add_run(line.lstrip('# ')); r.bold = True
        elif line.startswith('- '):
            p.add_run('• ' + line[2:])
        else:
            p.add_run(line.replace('**', ''))
    doc.save(os.path.join(B, out)); print('saved', out)


mk('cover_letter_AAP.md', 'Cover_letter_AAP.docx')
mk('highlights.txt', 'Highlights_AAP.docx')
mk('LEIA-ME_AAP.md', 'LEIA-ME_AAP.docx', 11)
